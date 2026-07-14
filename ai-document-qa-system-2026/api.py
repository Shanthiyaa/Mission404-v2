"""
api.py — FastAPI REST layer for ALE Knowledge Assistant
Wraps member1/extractor, member2/retriever, member2/vector_store.
Run with: python run_api.py

Place this file in the ROOT of ai-document-qa-system-2026/
(same level as config.py, member1/, member2/, member3/)
"""

from sqlalchemy.dialects import postgresql
from sqlalchemy.dialects import postgresql
from sqlalchemy.dialects import postgresql
from sqlalchemy.dialects import postgresql
import os
import sys
import re

try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except AttributeError:
    pass

import json
import time
import uuid
import shutil
import logging
import threading
from pathlib import Path
from datetime import datetime, timezone
from typing import Optional
from contextlib import asynccontextmanager

import jwt
import bcrypt

# ── Database & Models ─────────────────────────────────────────────────────────
from db import init_db, SessionLocal, User, Document, Conversation, Message, UserActivity, Notification

JWT_SECRET = os.getenv("JWT_SECRET", "SUPER_SECRET_SECURITY_KEY_ALE_ASSISTANT_2026")
JWT_ALGORITHM = "HS256"

def hash_password(password: str) -> str:
    salt = bcrypt.gensalt()
    return bcrypt.hashpw(password.encode('utf-8'), salt).decode('utf-8')

def verify_password(password: str, hashed_password: str) -> bool:
    try:
        return bcrypt.checkpw(password.encode('utf-8'), hashed_password.encode('utf-8'))
    except Exception:
        return False

def create_access_token(user_id: int, email: str, role: Optional[str] = None) -> str:
    payload = {
        "user_id": user_id,
        "email": email,
        "role": role,
        "exp": time.time() + 7 * 24 * 3600  # 7 days
    }
    token = jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)
    if isinstance(token, bytes):
        token = token.decode('utf-8')
    return token

# ── FastAPI ───────────────────────────────────────────────────────────────────
from fastapi import FastAPI, UploadFile, File, Form, BackgroundTasks, HTTPException, Depends, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse, HTMLResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from pydantic import BaseModel

# ── Shared config ─────────────────────────────────────────────────────────────
from config import (
    UPLOAD_DIR,
    FAISS_INDEX_DIR,
    CHUNKS_JSON_PATH,
    OLLAMA_MODEL,
    OLLAMA_BASE_URL,
    TOP_K_RESULTS,
    RAG_PROMPT_TEMPLATE,
    LLM_TEMPERATURE,
    API_HOST,
    API_PORT,
    CORS_ORIGINS,
)

# ── Logging ───────────────────────────────────────────────────────────────────
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger(__name__)

# ── Paths ─────────────────────────────────────────────────────────────────────
UPLOAD_PATH  = Path(UPLOAD_DIR)
FAISS_PATH   = Path(FAISS_INDEX_DIR)
CHUNKS_PATH  = Path(CHUNKS_JSON_PATH)

UPLOAD_PATH.mkdir(parents=True, exist_ok=True)
FAISS_PATH.mkdir(parents=True, exist_ok=True)
(Path("./data")).mkdir(parents=True, exist_ok=True)

# ── In-memory task tracker ────────────────────────────────────────────────────
_tasks: dict[str, dict] = {}

# ── Cached FAISS retriever ────────────────────────────────────────────────────
_retriever_cache = {"loaded": False}

# ── Stats counter ─────────────────────────────────────────────────────────────
_query_counter = {"total": 0, "confidence_sum": 0.0}

# ── Lock to prevent concurrent uploads from corrupting the FAISS index ────────
_faiss_lock = threading.Lock()

# ── DB session dependency ──────────────────────────────────────────────────────
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# ── Auth dependency ────────────────────────────────────────────────────────────
reusable_oauth2 = HTTPBearer(auto_error=False)

def get_current_user(
    credentials: Optional[HTTPAuthorizationCredentials] = Depends(reusable_oauth2),
    token: Optional[str] = Query(None)
):
    actual_token = None
    if credentials:
        actual_token = credentials.credentials
    elif token:
        actual_token = token

    if not actual_token:
        raise HTTPException(status_code=401, detail="Session expired or invalid authorization token.")

    try:
        payload = jwt.decode(actual_token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user_id = payload.get("user_id")
        if not user_id:
            raise HTTPException(status_code=401, detail="Invalid authorization token.")
            
        # Update last_activity (throttled to 30s)
        db = SessionLocal()
        try:
            user = db.query(User).filter(User.id == user_id).first()
            if user:
                now = datetime.utcnow()
                if not user.last_activity or (now - user.last_activity).total_seconds() > 30:
                    user.last_activity = now
                    db.commit()
        except Exception as e:
            log.warning(f"Failed to update last_activity: {e}")
        finally:
            db.close()
            
        return payload
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Session expired or invalid authorization token.")

def check_admin(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    user = db.query(User).filter(User.id == user_id).first()
    if not user or user.role != "Admin":
        raise HTTPException(status_code=403, detail="Access denied. Administrator privileges required.")
    return user

# ══════════════════════════════════════════════════════════════════════════════
#  Lifespan
# ══════════════════════════════════════════════════════════════════════════════

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Initialize SQL database tables
    init_db()
    _ensure_retriever()
    if _retriever_cache.get("loaded"):
        try:
            from member2.retriever import _get_embeddings
            _get_embeddings().embed_query("warmup")
            log.info("Embedding model warmed up.")
        except Exception as exc:
            log.warning(f"Embedding warmup skipped: {exc}")
    log.info("ALE Knowledge API started.")
    log.info(f"  Upload dir  : {UPLOAD_PATH.resolve()}")
    log.info(f"  FAISS index : {FAISS_PATH.resolve()}")
    log.info(f"  Ollama model: {OLLAMA_MODEL} @ {OLLAMA_BASE_URL}")
    log.info(f"  FAISS ready : {_retriever_cache['loaded']}")
    yield

# ── App ───────────────────────────────────────────────────────────────────────
app = FastAPI(title="ALE Knowledge API", version="1.0.0", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ══════════════════════════════════════════════════════════════════════════════
#  HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def _append_activity(user_id: int, text: str, color: str = "bg-purple-500") -> None:
    db = SessionLocal()
    try:
        evt = UserActivity(
            user_id=user_id,
            text=text,
            color=color,
            timestamp=datetime.utcnow()
        )
        db.add(evt)
        db.commit()
    except Exception as e:
        log.error(f"Failed to append activity in DB: {e}")
    finally:
        db.close()


def _create_notification(
    user_id: int, 
    type: str, 
    text: str, 
    link: str = None,
    title: str = None,
    target_conv_id: str = None,
    target_msg_id: int = None
) -> None:
    db = SessionLocal()
    try:
        notif = Notification(
            user_id=user_id,
            type=type,
            text=text,
            link=link,
            is_read=False,
            created_at=_utc_now(),
            title=title,
            target_conv_id=target_conv_id,
            target_msg_id=target_msg_id
        )
        db.add(notif)
        db.commit()
    except Exception as e:
        log.error(f"Failed to create notification: {e}")
    finally:
        db.close()


def _utc_now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _utc_iso(dt: datetime) -> str:
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def _format_time_ago(iso_str: str) -> str:
    try:
        dt = datetime.fromisoformat(iso_str.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        diff = int((datetime.now(timezone.utc) - dt.astimezone(timezone.utc)).total_seconds())
        diff = max(diff, 0)
        if diff < 60:
            return "Just now"
        if diff < 3600:
            minutes = diff // 60
            return f"{minutes} minute{'s' if minutes != 1 else ''} ago"
        if diff < 86400:
            hours = diff // 3600
            return f"{hours} hour{'s' if hours != 1 else ''} ago"
        days = diff // 86400
        return f"{days} day{'s' if days != 1 else ''} ago"
    except Exception:
        return "recently"


def _detect_category_from_filename(filename: str, fallback_category: str) -> str:
    name = Path(filename).stem.lower()
    normalized = re.sub(r"[_\-.]+", " ", name)
    normalized = re.sub(r"\s+", " ", normalized).strip()
    compact = re.sub(r"[^a-z0-9]", "", normalized)

    category_keywords = [
        ("user_guide", ["user guide", "user guides", "guide", "manual"], ["userguide", "userguides"]),
        ("release_note", ["release note", "release notes"], ["releasenote", "releasenotes"]),
        ("kcs", ["article", "articles", "knowledge", "kcs", "faq"], []),
        ("sqa", ["test case", "test cases", "testing", "validation"], ["testcase", "testcases"]),
    ]

    for category, phrases, compact_phrases in category_keywords:
        if any(phrase in normalized for phrase in phrases):
            return category
        if any(phrase in compact for phrase in compact_phrases):
            return category

    return fallback_category


def _get_pdf_page_count(filepath: Path) -> int:
    try:
        import re
        data = filepath.read_bytes()
        matches = re.findall(rb"/Type\s*/Page[^s]", data)
        return max(len(matches), 1)
    except Exception:
        return 1


def _get_page_count(filepath: Path) -> int:
    ext = filepath.suffix.lower()
    if ext == ".pdf":
        return _get_pdf_page_count(filepath)
    elif ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(str(filepath))
            return len(prs.slides)
        except Exception:
            return 1
    elif ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(filepath), read_only=True)
            count = len(wb.sheetnames)
            wb.close()
            return count
        except Exception:
            return 1
    return 1


def _human_size(bytes_: int) -> str:
    for unit in ["B", "KB", "MB", "GB"]:
        if bytes_ < 1024:
            return f"{bytes_:.1f} {unit}"
        bytes_ /= 1024
    return f"{bytes_:.1f} GB"


def _ensure_retriever():
    idx_file  = FAISS_PATH / "index.faiss"
    meta_file = FAISS_PATH / "metadata.json"
    _retriever_cache["loaded"] = idx_file.exists() and meta_file.exists()

GREETING_PATTERN = re.compile(
    r'^(hi|hello|hey|hy|yo|sup|good\s*(morning|afternoon|evening)|greetings|howdy|thanks|thank\s*you|ok|okay|bye|goodbye)[\s!.,?]*$',
    re.IGNORECASE
)

def _is_smalltalk(text: str) -> bool:
    return bool(GREETING_PATTERN.match(text.strip()))


def _get_unique_filename(target_path: Path) -> Path:
    if not target_path.exists():
        return target_path
    stem = target_path.stem
    suffix = target_path.suffix
    parent = target_path.parent
    counter = 1
    while True:
        new_path = parent / f"{stem}_{counter}{suffix}"
        if not new_path.exists():
            return new_path
        counter += 1


# ══════════════════════════════════════════════════════════════════════════════
#  BACKGROUND PIPELINE TASK
# ══════════════════════════════════════════════════════════════════════════════

def _run_pipeline(task_id: str, file_path: str, doc_entries: list[dict], user_id: int, is_zip: bool = False) -> None:
    t = _tasks[task_id]
    filename = Path(file_path).name
    SUPPORTED_EXTENSIONS_IN_ZIP = {".pdf", ".docx", ".pptx", ".xlsx", ".txt"}

    try:
        extracted_paths = []
        if is_zip:
            import zipfile
            t.update({"stage": "Extracting ZIP archive…", "progress": 5})
            log.info(f"[{task_id}] Extracting ZIP: {filename}")
            
            with zipfile.ZipFile(file_path) as z:
                zip_files = [
                    name for name in z.namelist()
                    if Path(name).suffix.lower() in SUPPORTED_EXTENSIONS_IN_ZIP
                    and not name.startswith("__MACOSX")
                    and not Path(name).name.startswith(".")
                ]
                
                for entry in doc_entries:
                    name = entry.get("zip_member", entry["name"])
                    target_name = entry["name"]
                    dest_file = UPLOAD_PATH / target_name
                    with open(dest_file, "wb") as f_sub:
                        f_sub.write(z.read(name))
                    extracted_paths.append(str(dest_file.resolve()))
                    
                    file_size = dest_file.stat().st_size
                    entry["size"] = _human_size(file_size)
                    entry["size_bytes"] = file_size
        else:
            extracted_paths.append(file_path)

        # Stage 1: Text extraction
        t.update({"stage": f"Extracting text from {len(extracted_paths)} document(s)…", "progress": 10})
        log.info(f"[{task_id}] Extracting text from: {extracted_paths}")

        from member1.extractor import process_documents
        documents = process_documents(extracted_paths)

        if not documents:
            raise ValueError("Extraction returned no content — check the document(s).")

        t.update({"stage": f"Chunked into {len(documents)} pieces", "progress": 40})
        log.info(f"[{task_id}] Chunks: {len(documents)}")

        # Stage 2: Embeddings + FAISS (with lock)
        t.update({"stage": f"Generating embeddings for {len(documents)} chunks...", "progress": 55})
        log.info(f"[{task_id}] Embedding…")

        with _faiss_lock:
            from member2.vector_store import add_to_faiss_index, remove_filenames_from_faiss_index

            new_chunks = [
                {
                    "text":        d.page_content,
                    "page":        d.metadata.get("page", d.metadata.get("chunk_index", 0)),
                    "slide":       d.metadata.get("slide"),
                    "worksheet":   d.metadata.get("worksheet"),
                    "source_file": d.metadata.get("source", filename),
                    "section":     d.metadata.get("section", ""),
                    "doc_type":    d.metadata.get("doc_type", "unknown"),
                    "is_table":    d.metadata.get("is_table", False),
                    "word_count":  d.metadata.get("word_count", 0),
                    "chunk_index": d.metadata.get("chunk_index", 0),
                }
                for d in documents
            ]
            log.info(f"[{task_id}] New chunks to embed: {len(new_chunks)}")

            meta_file = FAISS_PATH / "metadata.json"
            index_exists = meta_file.exists() and (FAISS_PATH / "index.faiss").exists()

            if index_exists:
                try:
                    # Remove any existing vectors for these filenames first
                    # This uses reconstruct_n — NO embedding regeneration
                    sub_filenames = [entry["name"] for entry in doc_entries]
                    remaining = remove_filenames_from_faiss_index(sub_filenames)
                    log.info(
                        f"[{task_id}] Cleaned overlapping filenames. "
                        f"{len(remaining)} existing chunks remain in index."
                    )
                except Exception as e:
                    log.warning(f"[{task_id}] Could not clean existing index: {e}. Continuing.")

            # Only embed the NEW chunks — never re-embed existing ones
            add_to_faiss_index(new_chunks)
            
            # Read back final metadata count for logging
            try:
                with open(meta_file, encoding="utf-8") as f:
                    merged_chunks = json.load(f)
            except Exception:
                merged_chunks = new_chunks

        t.update({"stage": "Indexing to FAISS…", "progress": 85})
        log.info(f"[{task_id}] FAISS index built with {len(merged_chunks)} total chunks.")

        # Stage 3: Register documents in SQL database
        db = SessionLocal()
        try:
            for entry in doc_entries:
                sub_filename = entry["name"]
                sub_chunks_count = sum(1 for d in documents if d.metadata.get("source") == sub_filename)
                sub_path = UPLOAD_PATH / sub_filename
                page_count = _get_page_count(sub_path)

                doc = db.query(Document).filter(Document.id == entry["id"]).first()
                if doc:
                    doc.status = "Indexed"
                    doc.chunks = sub_chunks_count
                    doc.pages = page_count
                    doc.size = entry["size"]
                    doc.size_bytes = entry["size_bytes"]
                else:
                    db_doc = Document(
                        id=entry["id"],
                        user_id=user_id,
                        name=entry["name"],
                        category=entry["category"],
                        size=entry["size"],
                        size_bytes=entry["size_bytes"],
                        pages=page_count,
                        status="Indexed",
                        task_id=task_id,
                        chunks=sub_chunks_count,
                        uploaded_at=datetime.utcnow()
                    )
                    db.add(db_doc)
            db.commit()
        except Exception as e:
            db.rollback()
            log.error(f"[{task_id}] Failed to save indexing stats to DB: {e}")
        finally:
            db.close()

        _retriever_cache["loaded"] = True
        
        try:
            from member2.retriever import clear_cache
            clear_cache()
        except ImportError:
            pass

        _append_activity(
            user_id,
            f"{filename} indexed successfully ({len(documents)} chunks)",
            color="bg-purple-500"
        )

        _create_notification(
            user_id,
            type="doc_processed",
            text=f'Document "{filename}" has been indexed successfully.',
            link=f"/knowledge-base?highlight={doc_entries[0]['id']}",
            title="Document Processing Completed",
            target_conv_id=doc_entries[0]['id']
        )

        t.update({
            "stage":    "Complete",
            "progress": 100,
            "done":     True,
            "chunks":   len(documents),
        })
        log.info(f"[{task_id}] ✅ Done.")

    except Exception as exc:
        log.error(f"[{task_id}] ✗ Pipeline failed: {exc}", exc_info=True)

        db = SessionLocal()
        try:
            for entry in doc_entries:
                doc = db.query(Document).filter(Document.id == entry["id"]).first()
                if doc:
                    doc.status = "Failed"
            db.commit()
        except Exception:
            db.rollback()
        finally:
            db.close()

        _append_activity(user_id, f"Failed to index {filename}: {exc}", color="bg-red-500")

        _create_notification(
            user_id,
            type="doc_failed",
            text=f'Failed to index "{filename}": {exc}',
            link=f"/knowledge-base?highlight={doc_entries[0]['id']}",
            title="Document Processing Failed",
            target_conv_id=doc_entries[0]['id']
        )

        t.update({
            "stage":    "Failed",
            "progress": 0,
            "done":     True,
            "error":    str(exc),
        })


# ── Auth Request Models ───────────────────────────────────────────────────────
class SignupRequest(BaseModel):
    name: str
    email: str
    department: str
    password: str

class LoginRequest(BaseModel):
    email: str
    password: str

class RegularLoginRequest(BaseModel):
    name: str
    email: str
    role: str

@app.post("/api/auth/signup")
async def signup(req: SignupRequest, db = Depends(get_db)):
    if len(req.password) < 8:
        raise HTTPException(status_code=400, detail="Password must be at least 8 characters.")
        
    existing = db.query(User).filter(User.email == req.email).first()
    if existing:
        raise HTTPException(status_code=400, detail="This email is already registered. Please sign in.")
    
    hashed = hash_password(req.password)
    user = User(
        username=req.name,
        email=req.email,
        hashed_password=hashed,
        department=req.department,
        role="Admin",  # Signup page is only for admins
        last_activity=datetime.utcnow()
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return {"success": True, "message": "Admin registered successfully."}

@app.post("/api/auth/login")
async def login(req: LoginRequest, db = Depends(get_db)):
    if len(req.password) < 8:
        raise HTTPException(status_code=400, detail="Password must be at least 8 characters.")
        
    user = db.query(User).filter(User.email == req.email).first()
    if not user or user.role != "Admin":
        raise HTTPException(status_code=401, detail="Incorrect email or password, or access denied.")
        
    if not verify_password(req.password, user.hashed_password):
        raise HTTPException(status_code=401, detail="Incorrect email or password.")
    
    user.last_activity = datetime.utcnow()
    db.commit()

    token = create_access_token(user.id, user.email, user.role)
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "name": user.username,
            "display_name": user.display_name or user.username,
            "email": user.email,
            "department": user.department,
            "profile_picture": user.profile_picture,
            "role": user.role
        }
    }

@app.post("/api/auth/login-regular")
async def login_regular(req: RegularLoginRequest, db = Depends(get_db)):
    if req.role == "Admin":
        raise HTTPException(status_code=400, detail="Admins must log in via the Admin Login page.")
        
    user = db.query(User).filter(User.email == req.email).first()
    if user:
        user.username = req.name
        user.role = req.role
        user.last_activity = datetime.utcnow()
    else:
        import secrets
        dummy_pass = secrets.token_hex(16)
        user = User(
            username=req.name,
            email=req.email,
            role=req.role,
            hashed_password=hash_password(dummy_pass),
            department="",
            last_activity=datetime.utcnow()
        )
        db.add(user)
    
    db.commit()
    db.refresh(user)

    token = create_access_token(user.id, user.email, user.role)
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "name": user.username,
            "display_name": user.display_name or user.username,
            "email": user.email,
            "department": user.department or "",
            "profile_picture": user.profile_picture,
            "role": user.role
        }
    }

@app.post("/api/auth/heartbeat")
async def auth_heartbeat(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    user = db.query(User).filter(User.id == user_id).first()
    if user:
        user.last_activity = datetime.utcnow()
        db.commit()
    return {"success": True}

@app.get("/api/auth/profile")
async def get_profile(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user = db.query(User).filter(User.id == current_user["user_id"]).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
    return {
        "name": user.username,
        "display_name": user.display_name or user.username,
        "email": user.email,
        "department": user.department or "",
        "profile_picture": user.profile_picture,
        "role": user.role
    }

class UpdateProfileRequest(BaseModel):
    username: Optional[str] = None
    display_name: Optional[str] = None
    email: Optional[str] = None
    profile_picture: Optional[str] = None
    current_password: Optional[str] = None
    new_password: Optional[str] = None

@app.post("/api/auth/profile/update")
async def update_profile(req: UpdateProfileRequest, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")

    if req.new_password:
        if not req.current_password:
            raise HTTPException(status_code=400, detail="Current password is required to update password.")
        if not verify_password(req.current_password, user.hashed_password):
            raise HTTPException(status_code=400, detail="Incorrect current password.")
        if len(req.new_password) < 8:
            raise HTTPException(status_code=400, detail="New password must be at least 8 characters.")
        user.hashed_password = hash_password(req.new_password)

    if req.email and req.email != user.email:
        taken = db.query(User).filter(User.email == req.email).first()
        if taken:
            raise HTTPException(status_code=400, detail="Email is already registered by another account.")
        user.email = req.email

    if req.username and req.username != user.username:
        user.username = req.username

    if req.display_name is not None:
        user.display_name = req.display_name

    if req.profile_picture is not None:
        user.profile_picture = req.profile_picture

    db.commit()
    db.refresh(user)

    token = create_access_token(user.id, user.email, user.role)
    return {
        "success": True,
        "access_token": token,
        "user": {
            "name": user.username,
            "display_name": user.display_name or user.username,
            "email": user.email,
            "department": user.department or "",
            "profile_picture": user.profile_picture,
            "role": user.role
        }
    }

# ── Admin User Management Endpoints ───────────────────────────────────────────
@app.get("/api/admin/users")
async def admin_list_users(admin_user: User = Depends(check_admin), db = Depends(get_db)):
    users = db.query(User).all()
    return [
        {
            "id": u.id,
            "username": u.username,
            "email": u.email,
            "role": u.role or "Normal User",
            "created_at": _utc_iso(u.created_at) if u.created_at else None
        }
        for u in users
    ]

@app.delete("/api/admin/users/{user_id}")
async def admin_delete_user(user_id: int, admin_user: User = Depends(check_admin), db = Depends(get_db)):
    if user_id == admin_user.id:
        raise HTTPException(status_code=400, detail="You cannot delete your own admin account.")
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
        
    user_docs = db.query(Document).filter(Document.user_id == user_id).all()
    for doc in user_docs:
        filename = doc.name
        file_path = UPLOAD_PATH / filename
        if file_path.exists():
            try:
                file_path.unlink()
            except Exception as e:
                log.error(f"Failed to unlink physical file {filename}: {e}")
        with _faiss_lock:
            try:
                from member2.vector_store import delete_from_faiss_index
                delete_from_faiss_index(filename)
            except Exception as e:
                log.error(f"Failed to delete {filename} from FAISS during user deletion: {e}")
                
    try:
        from member2.retriever import clear_cache
        clear_cache()
    except Exception:
        pass
    _ensure_retriever()
    
    db.delete(user)
    db.commit()
    return {"success": True}

# ── Notifications Endpoints ───────────────────────────────────────────────────
@app.get("/api/notifications")
async def list_notifications(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    notifs = db.query(Notification).filter(Notification.user_id == user_id).order_by(Notification.created_at.desc()).all()
    return [
        {
            "id": n.id,
            "type": n.type,
            "text": n.text,
            "link": n.link,
            "is_read": n.is_read,
            "created_at": _utc_iso(n.created_at),
            "time": _format_time_ago(_utc_iso(n.created_at)),
            "title": n.title,
            "target_conv_id": n.target_conv_id,
            "target_msg_id": n.target_msg_id
        }
        for n in notifs
    ]

@app.get("/api/notifications/unread-count")
async def unread_notifications_count(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    count = db.query(Notification).filter(Notification.user_id == user_id, Notification.is_read == False).count()
    return {"count": count}

@app.post("/api/notifications/{notif_id}/read")
async def mark_notification_read(notif_id: int, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    notif = db.query(Notification).filter(Notification.id == notif_id, Notification.user_id == user_id).first()
    if not notif:
        raise HTTPException(status_code=404, detail="Notification not found.")
    notif.is_read = True
    db.commit()
    return {"success": True}

@app.post("/api/notifications/read-all")
async def mark_all_notifications_read(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    db.query(Notification).filter(Notification.user_id == user_id).update({Notification.is_read: True}, synchronize_session=False)
    db.commit()
    return {"success": True}

@app.delete("/api/notifications/{notif_id}")
async def delete_single_notification(notif_id: int, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    notif = db.query(Notification).filter(Notification.id == notif_id, Notification.user_id == user_id).first()
    if not notif:
        raise HTTPException(status_code=404, detail="Notification not found.")
    db.delete(notif)
    db.commit()
    return {"success": True}

@app.delete("/api/notifications")
async def delete_all_notifications(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    db.query(Notification).filter(Notification.user_id == user_id).delete(synchronize_session=False)
    db.commit()
    return {"success": True}

# ── Conversations ─────────────────────────────────────────────────────────────
@app.get("/api/conversations")
async def list_conversations(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    convs = db.query(Conversation).filter(Conversation.user_id == user_id).order_by(Conversation.created_at.desc()).all()
    result = []
    for c in convs:
        msgs = db.query(Message).filter(Message.conversation_id == c.id).order_by(Message.id.asc()).all()
        messages_list = []
        for m in msgs:
            citations = json.loads(m.citations) if m.citations else None
            messages_list.append({
                "id": m.id,
                "role": m.role,
                "content": m.content,
                "citations": citations
            })
        result.append({
            "id": int(c.id) if str(c.id).isdigit() else c.id,
            "title": c.title,
            "messages": messages_list
        })
    return result

@app.delete("/api/conversations/{conv_id}")
async def delete_conversation(conv_id: str, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    conv = db.query(Conversation).filter(Conversation.id == conv_id, Conversation.user_id == user_id).first()
    if not conv:
        raise HTTPException(status_code=404, detail="Conversation not found.")
    db.delete(conv)
    db.commit()
    return {"success": True}

# ══════════════════════════════════════════════════════════════════════════════
#  UPLOAD  —  POST /api/upload
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/api/upload")
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    category: str    = Form("unknown"),
    current_user: dict = Depends(get_current_user),
    db = Depends(get_db)
):
    user_id = current_user["user_id"]
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
        
    # Calculate user's current storage space
    user_docs = db.query(Document).filter(Document.user_id == user_id).all()
    existing_storage = sum(doc.size_bytes for doc in user_docs)

    if user.role != "Admin":
        doc_count = len(user_docs)
        if doc_count >= 100:
            raise HTTPException(
                status_code=400,
                detail="You have reached the maximum document limit of 100. Please delete existing documents before uploading new documents."
            )
            
    filename_lower = file.filename.lower()
    is_zip = filename_lower.endswith(".zip")
    ext = Path(filename_lower).suffix.lower()
    
    SUPPORTED_EXTENSIONS_ALL = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".zip"}
    if ext not in SUPPORTED_EXTENSIONS_ALL:
        raise HTTPException(status_code=400, detail="Unsupported file format. Supported: PDF, DOCX, PPTX, XLSX, TXT, ZIP")

    task_id = str(uuid.uuid4())[:8]

    contents = await file.read()
    temp_dest = UPLOAD_PATH / f"temp_{uuid.uuid4()}{ext}"
    with open(temp_dest, "wb") as f:
        f.write(contents)

    # Perform duplicate check isolated to this user's documents
    try:
        from member1.extractor import calculate_file_hash
        target_hash = calculate_file_hash(str(temp_dest))
        user_docs = db.query(Document).filter(Document.user_id == user_id).all()
        user_doc_names = [d.name for d in user_docs]
        
        dup = False
        original = ""
        for name in user_doc_names:
            p = UPLOAD_PATH / name
            if p.exists() and p.is_file():
                if calculate_file_hash(str(p)) == target_hash:
                    dup = True
                    original = name
                    break
        
        if dup:
            temp_dest.unlink(missing_ok=True)
            _create_notification(
                user_id,
                type="duplicate_upload",
                text=f'Upload blocked: "{file.filename}" is a duplicate of "{original}".',
                link="/upload"
            )
            raise HTTPException(
                status_code=409,
                detail=f"Duplicate document detected. This file has the same content as '{original}' already in your knowledge base."
            )
    except HTTPException:
        raise
    except Exception as e:
        log.warning(f"Duplicate check failed (continuing): {e}")

    dest = _get_unique_filename(UPLOAD_PATH / file.filename)
    temp_dest.rename(dest)

    doc_entries = []
    
    if is_zip:
        SUPPORTED_EXTENSIONS_IN_ZIP = {".pdf", ".docx", ".pptx", ".xlsx", ".txt"}
        import zipfile
        try:
            with zipfile.ZipFile(dest) as z:
                zip_files = [
                    name for name in z.namelist()
                    if Path(name).suffix.lower() in SUPPORTED_EXTENSIONS_IN_ZIP
                    and not name.startswith("__MACOSX")
                    and not Path(name).name.startswith(".")
                ]
        except Exception as e:
            dest.unlink(missing_ok=True)
            raise HTTPException(status_code=400, detail=f"Invalid ZIP file: {e}")

        if not zip_files:
            dest.unlink(missing_ok=True)
            raise HTTPException(status_code=400, detail="No supported files (PDF, DOCX, PPTX, XLSX, TXT) found inside the ZIP archive.")

        # Load user's existing documents to check for duplicates
        user_docs = db.query(Document).filter(Document.user_id == user_id).all()
        user_doc_names = [d.name for d in user_docs]

        valid_zip_files = []
        duplicate_details = []
        
        temp_extract_dir = UPLOAD_PATH / f"temp_zip_{uuid.uuid4()}"
        temp_extract_dir.mkdir(parents=True, exist_ok=True)
        
        try:
            with zipfile.ZipFile(dest) as z:
                for name in zip_files:
                    sub_filename = Path(name).name
                    temp_sub_path = temp_extract_dir / sub_filename
                    
                    with open(temp_sub_path, "wb") as f_temp:
                        f_temp.write(z.read(name))
                    
                    from member1.extractor import calculate_file_hash
                    sub_hash = calculate_file_hash(str(temp_sub_path))
                    
                    dup = False
                    original = ""
                    for existing_name in user_doc_names:
                        p = UPLOAD_PATH / existing_name
                        if p.exists() and p.is_file():
                            if calculate_file_hash(str(p)) == sub_hash:
                                dup = True
                                original = existing_name
                                break
                    
                    if dup:
                        duplicate_details.append(f'"{sub_filename}" is a duplicate of "{original}"')
                        temp_sub_path.unlink()
                    else:
                        valid_zip_files.append((name, sub_filename, temp_sub_path))
        except Exception as e:
            shutil.rmtree(temp_extract_dir, ignore_errors=True)
            dest.unlink(missing_ok=True)
            raise HTTPException(status_code=500, detail=f"Error processing ZIP contents: {e}")

        if user.role != "Admin":
            total_zip_size = sum(temp_sub_path.stat().st_size for _, _, temp_sub_path in valid_zip_files)
            if existing_storage + total_zip_size > 512 * 1024 * 1024:
                shutil.rmtree(temp_extract_dir, ignore_errors=True)
                dest.unlink(missing_ok=True)
                raise HTTPException(
                    status_code=400,
                    detail="You have reached your 512 MB storage limit. Please delete documents or free up space to upload new files."
                )

            if doc_count + len(valid_zip_files) > 100:
                shutil.rmtree(temp_extract_dir, ignore_errors=True)
                dest.unlink(missing_ok=True)
                raise HTTPException(
                    status_code=400,
                    detail="You have reached the maximum document limit of 100. Please delete existing documents before uploading new documents."
                )

        if duplicate_details and not valid_zip_files:
            shutil.rmtree(temp_extract_dir, ignore_errors=True)
            dest.unlink(missing_ok=True)
            msg = "All files in the ZIP are duplicates: " + ", ".join(duplicate_details)
            _create_notification(
                user_id,
                type="duplicate_upload",
                text="Upload blocked: ZIP archive contains only duplicate files.",
                link="/upload"
            )
            raise HTTPException(status_code=409, detail=msg)

        # Notify the user about duplicates that were skipped, if any
        if duplicate_details:
            _create_notification(
                user_id,
                type="duplicate_upload",
                text=f"Skipped {len(duplicate_details)} duplicate documents from uploaded ZIP: {', '.join(duplicate_details[:2])}{'...' if len(duplicate_details) > 2 else ''}",
                link="/upload"
            )

        for name, sub_filename, temp_sub_path in valid_zip_files:
            target_path = _get_unique_filename(UPLOAD_PATH / sub_filename)
            shutil.move(str(temp_sub_path), str(target_path))
            detected_category = _detect_category_from_filename(sub_filename, category)
            
            doc_id = str(uuid.uuid4())[:12]
            sub_file_size = target_path.stat().st_size
            doc_entry = {
                "id":          doc_id,
                "name":        target_path.name,
                "zip_member":  name,
                "category":    detected_category,
                "size":        _human_size(sub_file_size),
                "size_bytes":  sub_file_size,
                "pages":       0,
                "status":      "Processing",
                "uploaded_at": datetime.utcnow().isoformat(),
                "task_id":     task_id,
                "chunks":      0,
            }
            doc_entries.append(doc_entry)
            
            db_doc = Document(
                id=doc_id,
                user_id=user_id,
                name=target_path.name,
                category=detected_category,
                size=_human_size(sub_file_size),
                size_bytes=sub_file_size,
                pages=0,
                status="Processing",
                task_id=task_id,
                chunks=0,
                uploaded_at=datetime.utcnow()
            )
            db.add(db_doc)
        db.commit()
        shutil.rmtree(temp_extract_dir, ignore_errors=True)
        
    else:
        file_size  = len(contents)
        if user.role != "Admin":
            if existing_storage + file_size > 512 * 1024 * 1024:
                dest.unlink(missing_ok=True)
                raise HTTPException(
                    status_code=400,
                    detail="You have reached your 512 MB storage limit. Please delete documents or free up space to upload new files."
                )
        page_count = _get_page_count(dest)
        doc_id  = str(uuid.uuid4())[:12]

        doc_entry = {
            "id":          doc_id,
            "name":        dest.name,
            "category":    category,
            "size":        _human_size(file_size),
            "size_bytes":  file_size,
            "pages":       page_count,
            "status":      "Processing",
            "uploaded_at": datetime.utcnow().isoformat(),
            "task_id":     task_id,
            "chunks":      0,
        }
        doc_entries.append(doc_entry)
        
        # Save Document DB model
        db_doc = Document(
            id=doc_id,
            user_id=user_id,
            name=dest.name,
            category=category,
            size=_human_size(file_size),
            size_bytes=file_size,
            pages=page_count,
            status="Processing",
            task_id=task_id,
            chunks=0,
            uploaded_at=datetime.utcnow()
        )
        db.add(db_doc)
        db.commit()

    _tasks[task_id] = {
        "task_id":  task_id,
        "filename": dest.name,
        "stage":    "Queued…",
        "progress": 0,
        "done":     False,
        "error":    None,
        "chunks":   0,
    }

    _append_activity(
        user_id,
        f"{dest.name} uploaded ({_human_size(len(contents))})",
        color="bg-purple-500"
    )

    background_tasks.add_task(_run_pipeline, task_id, str(dest), doc_entries, user_id, is_zip)

    return {
        "task_id":  task_id,
        "doc_id":   doc_entries[0]["id"],
        "filename": dest.name,
        "status":   "Processing",
        "message":  "Upload received. Pipeline started.",
    }


# ══════════════════════════════════════════════════════════════════════════════
#  UPLOAD STATUS  —  GET /api/upload/status/{task_id}
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/upload/status/{task_id}")
async def upload_status(task_id: str, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    if task_id in _tasks:
        return _tasks[task_id]

    # Check DB for documents with this task_id
    docs = db.query(Document).filter(Document.task_id == task_id).all()
    if not docs:
        raise HTTPException(status_code=404, detail="Task not found.")

    # If there are documents, check if they are completed
    statuses = [d.status for d in docs]
    if all(s == "Indexed" for s in statuses):
        return {
            "task_id":  task_id,
            "filename": docs[0].name if docs else "document",
            "stage":    "Complete",
            "progress": 100,
            "done":     True,
            "error":    None,
            "chunks":   sum(d.chunks for d in docs),
        }
    elif any(s == "Failed" for s in statuses):
        return {
            "task_id":  task_id,
            "filename": docs[0].name if docs else "document",
            "stage":    "Failed",
            "progress": 0,
            "done":     True,
            "error":    "Document processing failed",
            "chunks":   0,
        }
    else:
        # Task is not in _tasks but docs are still in "Processing" status in DB.
        # This indicates the server restarted or crashed during processing.
        # Mark them as "Failed" in the DB.
        for d in docs:
            d.status = "Failed"
        db.commit()
        return {
            "task_id":  task_id,
            "filename": docs[0].name if docs else "document",
            "stage":    "Failed",
            "progress": 0,
            "done":     True,
            "error":    "Server restarted, processing interrupted",
            "chunks":   0,
        }


# ══════════════════════════════════════════════════════════════════════════════
#  QUERY  —  POST /api/query
# ══════════════════════════════════════════════════════════════════════════════

class QueryRequest(BaseModel):
    question:   str
    session_id: Optional[str] = None
    top_k:      Optional[int] = TOP_K_RESULTS
    doc_files:  Optional[list[str]] = None


@app.post("/api/query")
async def query(req: QueryRequest, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    q = req.question.strip()
    if not q:
        raise HTTPException(status_code=400, detail="Question cannot be empty.")

    # 1. Get or create conversation record
    conv = db.query(Conversation).filter(Conversation.id == req.session_id, Conversation.user_id == user_id).first()
    if not conv:
        conv_count = db.query(Conversation).filter(Conversation.user_id == user_id).count()
        if conv_count >= 50:
            raise HTTPException(
                status_code=400,
                detail="You have reached the maximum limit. Please delete old conversations to continue."
            )
        conv = Conversation(id=req.session_id, user_id=user_id, title=q[:40])
        db.add(conv)
        db.commit()
    else:
        user_msg_count = db.query(Message).filter(Message.conversation_id == req.session_id, Message.role == "user").count()
        if user_msg_count >= 50:
            raise HTTPException(
                status_code=400,
                detail="You have reached the maximum limit. Please delete old conversations to continue."
            )

    # Save user message
    selected_docs_json = json.dumps(req.doc_files) if req.doc_files else "[]"
    user_msg = Message(
        conversation_id=req.session_id,
        role="user",
        content=q,
        selected_docs=selected_docs_json,
        timestamp=datetime.utcnow()
    )
    db.add(user_msg)
    db.commit()

    if _is_smalltalk(q):
        ans = "Hello! I'm your Ale Docbot. Ask me anything about your uploaded user guides, release notes, SQA test cases, and KCS articles."
        assist_msg = Message(
            conversation_id=req.session_id,
            role="assistant",
            content=ans,
            citations="[]",
            selected_docs=selected_docs_json,
            timestamp=datetime.utcnow()
        )
        db.add(assist_msg)
        db.commit()
        db.refresh(user_msg)
        db.refresh(assist_msg)
        return {
            "answer": ans,
            "citations": [],
            "confidence": 100,
            "session_id": req.session_id,
            "message_id": assist_msg.id,
            "user_message_id": user_msg.id,
        }

    # ── Smart History Lookup (Semantic Cache) ─────────────────────────────────
    def _normalize_question(text: str) -> str:
        return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text.lower())).strip()

    def _same_selected_docs(prev_docs_json: str | None) -> bool:
        try:
            prev_docs = sorted(json.loads(prev_docs_json or "[]"))
            curr_docs = sorted(json.loads(selected_docs_json or "[]"))
            return prev_docs == curr_docs
        except Exception:
            return False

    def _assistant_after(user_message: Message) -> Optional[Message]:
        return db.query(Message).filter(
            Message.conversation_id == user_message.conversation_id,
            Message.id > user_message.id,
            Message.role == "assistant"
        ).order_by(Message.id.asc()).first()

    prev_user_msgs = db.query(Message).join(Conversation).filter(
        Conversation.user_id == user_id,
        Message.role == "user",
        Message.id < user_msg.id
    ).order_by(Message.id.desc()).limit(100).all()

    current_norm = _normalize_question(q)

    for prev_msg in prev_user_msgs:
        if _normalize_question(prev_msg.content) == current_norm and _same_selected_docs(prev_msg.selected_docs):
            matched_assist_msg = _assistant_after(prev_msg)
            if matched_assist_msg:
                log.info(f"History Cache HIT: exact question match for '{q}'")
                citations = json.loads(matched_assist_msg.citations) if matched_assist_msg.citations else []

                assist_msg = Message(
                    conversation_id=req.session_id,
                    role="assistant",
                    content=matched_assist_msg.content,
                    citations=matched_assist_msg.citations,
                    selected_docs=selected_docs_json,
                    timestamp=datetime.utcnow()
                )
                db.add(assist_msg)
                db.commit()
                db.refresh(user_msg)
                db.refresh(assist_msg)

                return {
                    "answer": matched_assist_msg.content,
                    "citations": citations,
                    "confidence": 100,
                    "session_id": req.session_id,
                    "cached": True,
                    "cache_type": "exact_history",
                    "message_id": assist_msg.id,
                    "user_message_id": user_msg.id,
                }

    semantic_candidates = [
        m for m in prev_user_msgs
        if _same_selected_docs(m.selected_docs) and _assistant_after(m)
    ]

    if semantic_candidates:
        try:
            from member2.retriever import _get_embeddings
            import numpy as np
            embeddings = _get_embeddings()
            q_vec = np.array(embeddings.embed_query(q), dtype="float32")
            history_questions = [m.content for m in semantic_candidates]
            hist_vecs = np.array(embeddings.embed_documents(history_questions), dtype="float32")
            
            q_norm = np.linalg.norm(q_vec)
            best_sim = 0.0
            best_msg = None
            if q_norm > 0:
                for hist_msg, h_vec in zip(semantic_candidates, hist_vecs):
                    h_norm = np.linalg.norm(h_vec)
                    if h_norm > 0:
                        sim = np.dot(q_vec, h_vec) / (q_norm * h_norm)
                        if sim > best_sim:
                            best_sim = sim
                            best_msg = hist_msg
            
            if best_sim >= 0.92 and best_msg:
                matched_assist_msg = _assistant_after(best_msg)
                if matched_assist_msg:
                    log.info(f"Smart Cache HIT: Matched question '{q}' with '{best_msg.content}' (score: {best_sim:.3f})")
                    citations = json.loads(matched_assist_msg.citations) if matched_assist_msg.citations else []

                    assist_msg = Message(
                        conversation_id=req.session_id,
                        role="assistant",
                        content=matched_assist_msg.content,
                        citations=matched_assist_msg.citations,
                        selected_docs=selected_docs_json,
                        timestamp=datetime.utcnow()
                    )
                    db.add(assist_msg)
                    db.commit()
                    db.refresh(user_msg)
                    db.refresh(assist_msg)
                    _create_notification(
                        user_id,
                        type="ai_answer",
                        text=f'Your answer for "{q[:40]}..." is ready.',
                        link=f"/chat?id={req.session_id}&msgId={assist_msg.id}",
                        title="Query Answer Completed",
                        target_conv_id=req.session_id,
                        target_msg_id=assist_msg.id
                    )

                    return {
                        "answer": matched_assist_msg.content,
                        "citations": citations,
                        "confidence": 100,
                        "session_id": req.session_id,
                        "cached": True,
                        "cache_type": "semantic_history",
                        "cache_similarity": round(float(best_sim), 3),
                        "message_id": assist_msg.id,
                        "user_message_id": user_msg.id,
                    }
        except Exception as e:
            log.warning(f"Smart History Lookup failed: {e}")

    # Fall back to RAG Pipeline
    if not _retriever_cache.get("loaded"):
        _ensure_retriever()
        if not _retriever_cache.get("loaded"):
            err_ans = "No documents have been indexed yet. Please upload a PDF first."
            assist_msg = Message(
                conversation_id=req.session_id,
                role="assistant",
                content=err_ans,
                citations="[]",
                selected_docs=selected_docs_json,
                timestamp=datetime.utcnow()
            )
            db.add(assist_msg)
            db.commit()
            raise HTTPException(
                status_code=503,
                detail="No documents indexed yet. Please upload a PDF first."
            )

    # USER-ISOLATED RETRIEVAL
    user_docs = db.query(Document).filter(
        Document.user_id == user_id,
        Document.status == "Indexed"
    ).all()
    user_doc_filenames = [d.name for d in user_docs]

    if req.doc_files:
        allowed_doc_files = [f for f in req.doc_files if f in user_doc_filenames]
    else:
        allowed_doc_files = user_doc_filenames

    if not allowed_doc_files:
        ans = "Please select at least one document before asking a question."
        assist_msg = Message(
            conversation_id=req.session_id,
            role="assistant",
            content=ans,
            citations="[]",
            selected_docs=selected_docs_json,
            timestamp=datetime.utcnow()
        )
        db.add(assist_msg)
        db.commit()
        return {
            "answer": ans,
            "citations": [],
            "confidence": 0,
            "session_id": req.session_id,
        }

    try:
        from member2.retriever import retrieve
        chunks = retrieve(q, k=req.top_k, doc_files=allowed_doc_files)
    except FileNotFoundError:
        raise HTTPException(
            status_code=503,
            detail="FAISS index not found. Please upload and index a document first."
        )
    except Exception as exc:
        log.error(f"Retrieval error: {exc}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Retrieval failed: {exc}")

    if not chunks:
        ans = "I couldn't find relevant content for your question in the uploaded documents."
        assist_msg = Message(
            conversation_id=req.session_id,
            role="assistant",
            content=ans,
            citations="[]",
            selected_docs=selected_docs_json,
            timestamp=datetime.utcnow()
        )
        db.add(assist_msg)
        db.commit()
        return {
            "answer":     ans,
            "citations":  [],
            "confidence": 0,
            "session_id": req.session_id,
        }

    # Build context for LLM
    context = "\n\n".join(
        f"[Source: {c['source_file']}, page {c['page']}]\n{c['text']}"
        for c in chunks
    )
    prompt = RAG_PROMPT_TEMPLATE.format(context=context, question=q)

    # Call Ollama
    try:
        import ollama
        response = ollama.chat(
            model=OLLAMA_MODEL,
            messages=[{"role": "user", "content": prompt}],
            options={
                "num_predict": 1024,
                "temperature": LLM_TEMPERATURE,
            },
        )
        answer = response.message.content.strip()
    except Exception as exc:
        log.error(f"Ollama error: {exc}", exc_info=True)
        raise HTTPException(
            status_code=503,
            detail=(
                f"LLM unavailable: {exc}. "
                "Make sure Ollama is running: `ollama serve` and model is pulled: `ollama pull llama3.2`"
            )
        )

    avg_score  = sum(c.get("score", 0) for c in chunks) / len(chunks)
    confidence = min(int(avg_score * 100), 99)

    _query_counter["total"]          += 1
    _query_counter["confidence_sum"] += confidence
    
    _append_activity(
        user_id,
        f"Query answered: \"{q[:60]}{'…' if len(q) > 60 else ''}\"",
        color="bg-green-500"
    )

    citations = []
    for i, c in enumerate(chunks):
        section   = c.get("section", "")
        full_text = c["text"]
        doc_type  = c.get("doc_type", "unknown")
        pdf_anchor = f"#page={c.get('page', 1)}" if doc_type == "pdf" else ""
        
        citations.append({
            "source_file":    c["source_file"],
            "page":           c.get("page", 1),
            "slide":          c.get("slide"),
            "worksheet":      c.get("worksheet"),
            "section":        section,
            "score":          round(c.get("score", 0), 3),
            "text":           full_text,
            "text_preview":   full_text[:300],
            "confidence":     min(int(c.get("score", 0) * 100), 99),
            "is_table":       c.get("is_table", False),
            "citation_label": f"Source {i+1}",
            "pdf_anchor":     pdf_anchor,
        })

    # Save assistant message in DB
    citations_json = json.dumps(citations)
    assist_msg = Message(
        conversation_id=req.session_id,
        role="assistant",
        content=answer,
        citations=citations_json,
        selected_docs=selected_docs_json,
        timestamp=datetime.utcnow()
    )
    try:
        db.add(assist_msg)
        db.commit()
        db.refresh(user_msg)
        db.refresh(assist_msg)
    except Exception as e:
        print("DATABASE ERROR:", repr(e))
        db.rollback()
        raise
    
    _create_notification(
        user_id,
        type="ai_answer",
        text=f'Your answer for "{q[:40]}..." is ready.',
        link=f"/chat?id={req.session_id}&msgId={assist_msg.id}",
        title="Query Answer Completed",
        target_conv_id=req.session_id,
        target_msg_id=assist_msg.id
    )

    return {
        "answer":     answer,
        "citations":  citations,
        "confidence": confidence,
        "session_id": req.session_id,
        "message_id": assist_msg.id,
        "user_message_id": user_msg.id,
    }


# ══════════════════════════════════════════════════════════════════════════════
#  DOCUMENTS  —  GET /api/documents
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/documents")
async def list_documents(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    docs = db.query(Document).filter(Document.user_id == user_id).order_by(Document.uploaded_at.desc()).all()

    result = []
    for d in docs:
        result.append({
            "id":          d.id,
            "name":        d.name,
            "category":    d.category,
            "size":        d.size,
            "size_bytes":  d.size_bytes,
            "pages":       d.pages,
            "status":      d.status,
            "uploaded_at": d.uploaded_at.isoformat(),
            "date":        _format_time_ago(d.uploaded_at.isoformat()),
            "chunks":      d.chunks,
            "task_id":     d.task_id,
        })
    return result


# ══════════════════════════════════════════════════════════════════════════════
#  VIEW DOCUMENT  —  GET /api/documents/{filename}/view
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/documents/{filename}/view")
async def view_document(filename: str, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    doc = db.query(Document).filter(Document.name == filename, Document.user_id == user_id).first()
    if not doc:
        raise HTTPException(status_code=403, detail="Access denied.")

    file_path = UPLOAD_PATH / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail=f"File '{filename}' not found.")
    
    ext = file_path.suffix.lower()
    media_type = "application/octet-stream"
    if ext == ".pdf":
        media_type = "application/pdf"
    elif ext == ".docx":
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif ext == ".pptx":
        media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif ext == ".xlsx":
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif ext in (".txt", ".log"):
        media_type = "text/plain"

    return FileResponse(
        path=str(file_path),
        media_type=media_type,
        headers={"Content-Disposition": f"inline; filename=\"{filename}\""}
    )


# ══════════════════════════════════════════════════════════════════════════════
#  DELETE DOCUMENT  —  DELETE /api/documents/{doc_id}
# ══════════════════════════════════════════════════════════════════════════════

@app.delete("/api/documents/{doc_id}")
async def delete_document(doc_id: str, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    target = db.query(Document).filter(Document.id == doc_id, Document.user_id == user_id).first()

    if not target:
        raise HTTPException(status_code=404, detail="Document not found.")

    filename = target.name
    db.delete(target)
    db.commit()

    file_path = UPLOAD_PATH / filename
    if file_path.exists():
        file_path.unlink()

    with _faiss_lock:
        try:
            from member2.vector_store import delete_from_faiss_index
            delete_from_faiss_index(filename)
            try:
                from member2.retriever import clear_cache
                clear_cache()
            except ImportError:
                pass
            _ensure_retriever()
        except Exception as exc:
            log.error(f"FAISS deletion failed: {exc}")

    _append_activity(user_id, f"{filename} deleted from knowledge base", color="bg-red-500")
    _create_notification(
        user_id,
        type="doc_deleted",
        text=f'"{filename}" has been successfully deleted.',
        link="/knowledge-base"
    )
    return {"success": True, "deleted": filename}


# ══════════════════════════════════════════════════════════════════════════════
#  STATS  —  GET /api/stats
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/stats")
async def get_stats(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    
    docs = db.query(Document).filter(Document.user_id == user_id).all()
    total_docs = len(docs)
    indexed_docs = sum(1 for d in docs if d.status == "Indexed")
    
    total_q = db.query(Message).join(Conversation).filter(
        Conversation.user_id == user_id,
        Message.role == "user"
    ).count()

    # Active Users
    from datetime import datetime, timedelta
    active_threshold = datetime.utcnow() - timedelta(minutes=5)
    active_users_count = db.query(User).filter(User.last_activity >= active_threshold).count()
    active_users_count = max(active_users_count, 1)

    # 1. Documents uploaded over time
    from collections import defaultdict
    upload_history = defaultdict(int)
    for d in docs:
        date_str = d.uploaded_at.strftime("%Y-%m-%d")
        upload_history[date_str] += 1
    sorted_uploads = [{"date": k, "count": v} for k, v in sorted(upload_history.items())]

    # 2. Queries per day
    user_queries = db.query(Message).join(Conversation).filter(
        Conversation.user_id == user_id,
        Message.role == "user"
    ).all()
    query_history = defaultdict(int)
    for mq in user_queries:
        date_str = mq.timestamp.strftime("%Y-%m-%d")
        query_history[date_str] += 1
    sorted_queries = [{"date": k, "count": v} for k, v in sorted(query_history.items())]

    # 3. Document categories
    category_counts = defaultdict(int)
    for d in docs:
        category_counts[d.category] += 1
    cat_data = [{"category": k, "count": v} for k, v in category_counts.items()]

    # 4. Processing status
    status_counts = defaultdict(int)
    for d in docs:
        status_counts[d.status] += 1
    status_data = [{"status": k, "count": v} for k, v in status_counts.items()]

    return {
        "total_documents":   total_docs,
        "indexed_documents": indexed_docs,
        "total_queries":     total_q,
        "avg_confidence":    0,
        "active_users":      active_users_count,
        "faiss_ready":       _retriever_cache.get("loaded", False),
        "charts": {
            "documents_over_time": sorted_uploads,
            "queries_per_day": sorted_queries,
            "document_categories": cat_data,
            "processing_status": status_data
        }
    }


# ══════════════════════════════════════════════════════════════════════════════
#  STORAGE  —  GET /api/storage/summary  &  POST /api/storage/free-up-space
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/storage/summary")
async def get_storage_summary(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
    
    docs = db.query(Document).filter(Document.user_id == user_id).all()
    used_bytes = sum(doc.size_bytes for doc in docs)
    limit_bytes = 512 * 1024 * 1024  # 512 MB
    
    return {
        "used_bytes": used_bytes,
        "limit_bytes": limit_bytes,
        "role": user.role or "Normal User"
    }


@app.post("/api/storage/free-up-space")
async def free_up_space(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    
    # 1. Get all documents belonging to this user
    docs = db.query(Document).filter(Document.user_id == user_id).all()
    if not docs:
        return {"success": True, "message": "Storage is already empty."}

    filenames = [d.name for d in docs]
    
    # 2. Delete rows from DB
    for doc in docs:
        db.delete(doc)
    db.commit()

    # 3. Delete physical files from UPLOAD_PATH
    for filename in filenames:
        file_path = UPLOAD_PATH / filename
        if file_path.exists():
            try:
                file_path.unlink()
            except Exception as e:
                log.error(f"Failed to unlink physical file {filename} during free-up-space: {e}")

    # 4. Remove files from FAISS index in a thread-safe way
    with _faiss_lock:
        try:
            from member2.vector_store import remove_filenames_from_faiss_index
            remove_filenames_from_faiss_index(filenames)
            try:
                from member2.retriever import clear_cache
                clear_cache()
            except ImportError:
                pass
            _ensure_retriever()
        except Exception as exc:
            log.error(f"FAISS batch deletion during free-up-space failed: {exc}")

    # 5. Log activity and send notification
    _append_activity(user_id, "Cleared all files from knowledge base (Free Up Space)", color="bg-red-500")
    _create_notification(
        user_id,
        type="doc_deleted",
        text="All documents have been deleted from your knowledge base to free up space.",
        link="/knowledge-base"
    )

    return {"success": True, "message": "Storage cleared successfully."}


# ══════════════════════════════════════════════════════════════════════════════
#  ACTIVITY  —  GET /api/activity
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/activity")
async def get_activity(current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    events = db.query(UserActivity).filter(UserActivity.user_id == user_id).order_by(UserActivity.timestamp.desc()).limit(10).all()
    return [
        {
            "text":  e.text,
            "time":  _format_time_ago(e.timestamp.isoformat()),
            "color": e.color,
        }
        for e in events
    ]


# ══════════════════════════════════════════════════════════════════════════════
#  HEALTH  —  GET /api/health
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/health")
async def health():
    return {
        "status":      "ok",
        "faiss_ready": _retriever_cache.get("loaded", False),
        "ollama_url":  OLLAMA_BASE_URL,
        "model":       OLLAMA_MODEL,
    }


@app.get("/api/documents/{filename}/view-html")
async def view_document_html(filename: str, current_user: dict = Depends(get_current_user), db = Depends(get_db)):
    user_id = current_user["user_id"]
    doc = db.query(Document).filter(Document.name == filename, Document.user_id == user_id).first()
    if not doc:
        raise HTTPException(status_code=403, detail="Access denied.")

    file_path = UPLOAD_PATH / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail=f"File '{filename}' not found.")
    
    ext = file_path.suffix.lower()
    try:
        if ext == ".docx":
            from member1.extractor import docx_to_html
            html_content = docx_to_html(file_path)
        elif ext == ".pptx":
            from member1.extractor import pptx_to_html
            html_content = pptx_to_html(file_path)
        elif ext == ".xlsx":
            from member1.extractor import xlsx_to_html
            html_content = xlsx_to_html(file_path)
        elif ext in (".txt", ".log"):
            content = file_path.read_text(encoding="utf-8", errors="replace")
            import html
            content_escaped = html.escape(content)
            html_content = f"<pre style='font-family: monospace; white-space: pre-wrap; padding: 20px; max-width: 900px; margin: 0 auto;'>{content_escaped}</pre>"
        else:
            raise HTTPException(status_code=400, detail="HTML preview not supported for this file type.")
        
        return HTMLResponse(content=html_content, status_code=200)
    except Exception as e:
        log.error(f"Failed to generate HTML preview for {filename}: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Preview generation failed: {e}")


# ══════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("api:app", host=API_HOST, port=API_PORT, reload=True)
